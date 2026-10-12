"""Render selected XLSX cells for layout review without claiming native Excel rendering."""
from pathlib import Path
import json
from PIL import Image,ImageDraw,ImageFont
import openpyxl
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parent
path=ROOT/'output'/'ECS_Fargate_HealthCheck_詳細設計_20261012.xlsx'
review=ROOT/'_review'
review.mkdir(exist_ok=True)
wb=openpyxl.load_workbook(path,data_only=True)


def font(size,bold=False):
    file='meiryob.ttc' if bold else 'meiryo.ttc'
    return ImageFont.truetype(str(Path(r'C:\Windows\Fonts')/file),round(size*4/3),index=2)


def wrap(text,face,width):
    lines=[]
    for para in str(text).split('\n'):
        current=''
        for char in para:
            if current and face.getlength(current+char)>width:
                lines.append(current)
                current=''
            current+=char
        lines.append(current)
    return lines


def rgb(color,default):
    if color is not None and color.type=='rgb':
        return '#'+color.rgb[-6:]
    return default


def render(name,rows,filename):
    ws=wb[name]
    col_widths=[round((ws.column_dimensions[get_column_letter(i)].width or 13)*7) for i in range(1,ws.max_column+1)]
    xs=[0]
    for w in col_widths: xs.append(xs[-1]+w)
    row_heights=[round((ws.row_dimensions[r].height or 15)*4/3) for r in rows]
    image=Image.new('RGB',(xs[-1]+2,sum(row_heights)+62),'white')
    draw=ImageDraw.Draw(image)
    label='セル配置からの確認用描画（Excel本体のスクリーンショットではありません）'
    draw.text((8,12),label,font=font(10),fill='#555555')
    y=42
    overflow=[]
    for row,height in zip(rows,row_heights):
        for col in range(1,ws.max_column+1):
            cell=ws.cell(row,col)
            if isinstance(cell,openpyxl.cell.cell.MergedCell):
                continue
            right=col
            for merged in ws.merged_cells.ranges:
                if cell.coordinate in merged:
                    right=merged.max_col
                    break
            x0,x1=xs[col-1],xs[right]
            fill=rgb(cell.fill.fgColor,'#FFFFFF') if cell.fill.patternType=='solid' else '#FFFFFF'
            draw.rectangle((x0,y,x1,y+height),fill=fill,outline='#D0D0D0')
            if cell.value is None: continue
            face=font(cell.font.sz or 11,cell.font.b)
            color=rgb(cell.font.color,'#242424')
            text=str(cell.value)
            if isinstance(cell.value,(int,float)) and '秒' in cell.number_format: text=f'{cell.value} 秒'
            lines=wrap(text,face,x1-x0-14)
            step=round((cell.font.sz or 11)*1.58)
            if len(lines)*step+10>height:
                overflow.append(f'{name}!{cell.coordinate}: {len(lines)*step+10}>{height}')
            yy=y+6
            for line in lines:
                xx=x0+7
                if cell.alignment.horizontal=='center': xx=x0+(x1-x0-face.getlength(line))/2
                draw.text((xx,yy),line,font=face,fill=color)
                yy+=step
        y+=height
    image.save(review/filename)
    return overflow


overflow=[]
overflow+=render('00_ガイド',range(1,16),'guide.png')
overflow+=render('06_バッファ内訳',range(1,17),'buffer_1.png')
overflow+=render('06_バッファ内訳',[1,2,3,4,5]+list(range(25,33)),'buffer_grace.png')
overflow+=render('07_パターン比較',range(1,13),'patterns.png')
overflow+=render('21_秒数試算',[1,2,3,4,5]+list(range(62,77)),'calculator.png')
result={'method':'XLSXセル値・行高・フォントからのPillow描画。Excel本体ではない。','font':font(11).getname(),'overflow':overflow}
(review/'layout_review.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
