"""Validate report content, numeric examples and spreadsheet structure locally."""
from __future__ import annotations
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import openpyxl
import yaml
import report_data as d

ROOT=Path(__file__).resolve().parent
OUT=ROOT/"output"
REPORT=OUT/"ECS_Fargate_HealthCheck_詳細設計_20261012.xlsx"
MARKDOWN=REPORT.with_suffix(".md")


def validate():
    checks=[]
    def require(ok,message):
        if not ok:
            raise AssertionError(message)
        checks.append(message)

    initial_expected={
        "P0":(180,195,225,285,True),
        "P1":(218,233,263,323,True),
        "P2":(138,153,183,243,True),
        "P3":(236,251,281,341,True),
        "P4":(386,401,431,491,False),
        "P5":(218,233,263,323,True),
        "P6":(218,233,263,323,True),
    }
    for p in d.PATTERNS:
        m=d.calculate(p)
        actual=(m['endpoint'],m['f_h'],m['service'],m['required'],m['dependency_ok'])
        require(actual==initial_expected[p['id']],f"{p['id']}: 独立した既知の計算結果と一致")
        if p["id"]=="P4":
            continue
        definition=d.task_fragment(p)
        json.loads(json.dumps(definition))
        names={c['name'] for c in definition['containerDefinitions']}
        require(len(names)==4,f"{p['id']}: 4つの固有コンテナ")
        edges={}
        for c in definition['containerDefinitions']:
            h=c['healthCheck']
            require(c['essential'] is True,f"{p['id']}/{c['name']}: essential=true")
            require(0<=h['startPeriod']<=300 and 5<=h['interval']<=300 and 2<=h['timeout']<=60 and 1<=h['retries']<=10,f"{p['id']}/{c['name']}: health値域内")
            require(h['command'][0] in ['CMD','CMD-SHELL'],f"{p['id']}/{c['name']}: command形式")
            if 'startTimeout' in c:
                require(2<=c['startTimeout']<=120,f"{p['id']}/{c['name']}: startTimeout値域内")
            edges[c['name']]=[]
            for dependency in c.get('dependsOn',[]):
                require(dependency['containerName'] in names and dependency['condition'] in ['START','HEALTHY'],f"{p['id']}/{c['name']}: 有効な依存先")
                edges[c['name']].append(dependency['containerName'])
        def visit(name,path):
            if name in path:
                raise AssertionError(f"循環依存: {path} -> {name}")
            for target in edges[name]:
                visit(target,path+[name])
        for name in names:
            visit(name,[])
        checks.append(f"{p['id']}: 循環依存なし")

    p1=d.PATTERNS[1]
    baseline=d.calculate(p1)
    delayed_c=d.calculate(p1,{'c_process':400})
    require(delayed_c['service']>baseline['service'],"CWが最遅ならMAXに反映される")
    fast_front=d.calculate(p1,{'f_max':1,'f_var':0,'f_pre':0,'f_post':0})
    require(fast_front['endpoint']==baseline['b_ready'],"フロントが早くてもバック準備前に複合readyにしない")
    slow_front=d.calculate(p1,{'f_max':260,'f_var':0})
    require(slow_front['endpoint']==280,"フロント自身が遅い場合もMAXに反映")
    delay_adot=d.calculate(p1,{'a_config':200})
    require(delay_adot['dependency_ok'] is False,"ADOTが120秒を超える場合の依存不成立")
    later_grace=d.calculate(p1,{'grace_origin':38})
    require(later_grace['required']==285,"猶予開始差を引き、図の時計と混同しない")
    require(d.calculate(d.PATTERNS[3],{'b_max':100})['dependency_ok'] is False,"P3バック遅延時に120秒制約を検出")
    require(baseline['f_h']==233 and baseline['f_local']==233,"P1の並列開始からのフロント猶予根拠")

    md=MARKDOWN.read_text(encoding='utf-8')
    require('\ufffd' not in md,"Markdownに置換文字なし")
    require(md.count('```')%2==0,"Markdownコードフェンスが閉じる")
    parsed=[]
    for block in re.findall(r'```json\n(.*?)\n```',md,re.S):
        parsed.append(json.loads(block))
    require(len(parsed)==8,"Markdown内の8つのJSON断片を構文検査")
    for block in re.findall(r'```yaml\n(.*?)\n```',md,re.S):
        obj=yaml.safe_load(block)
        require('health_check' in obj['service']['extensions'],"ADOT health extensionが有効化される断片")
    for source in d.SOURCES:
        require(source in md,f"{source}: 出典記載あり")
    for sec in d.SECTIONS:
        require(sec['title'] in md,f"{sec['sheet']}: Markdownと内容共有")
    for path in [ROOT/'report_data.py',ROOT/'build_report.py',ROOT/'validate_report.py',MARKDOWN]:
        text=path.read_text(encoding='utf-8')
        require(all(line==line.rstrip() for line in text.splitlines()),f"{path.name}: 末尾空白なし")
        require(text.endswith('\n'),f"{path.name}: 最終改行あり")

    wb=openpyxl.load_workbook(REPORT,data_only=False)
    cached=openpyxl.load_workbook(REPORT,data_only=True)
    manifest=json.loads((OUT/'calculation_manifest.json').read_text(encoding='utf-8'))
    require(len(wb.sheetnames)==25,"Excelは25シート")
    require(wb.sheetnames==manifest['sheets'],"シート索引と実際の順序が一致")
    errors=[]; fonts=set(); formulas=0
    for ws in wb.worksheets:
        require(ws.freeze_panes is not None,f"{ws.title}: 固定枠あり")
        require(bool(ws.print_area),f"{ws.title}: 印刷範囲あり")
        require(ws.sheet_view.showGridLines is False,f"{ws.title}: グリッド非表示")
        for row in ws:
            for cell in row:
                if cell.value is None:
                    continue
                fonts.add(cell.font.name)
                if cell.data_type=='f':
                    formulas+=1
                    cv=cached[ws.title][cell.coordinate]
                    if cv.data_type=='e' or cv.value is None:
                        errors.append((ws.title,cell.coordinate,cv.value))
                if cell.data_type=='e':
                    errors.append((ws.title,cell.coordinate,cell.value))
                if isinstance(cell.value,str):
                    require('\ufffd' not in cell.value,f"{ws.title}!{cell.coordinate}: 文字データ正常")
    require(fonts=={'Meiryo UI'},"全ての値のあるセルがMeiryo UI")
    require(not errors,"数式エラー・未格納の初期結果なし")
    require(formulas==33,"秒数試算に33個の数式")
    for key,expected in [('a_h',38),('b_local',195),('b_ready',218),('f_wait',58),('f_h',233),('normal',263),('spare',60),('required',323),('suggested',360),('margin',37),('g300',-23),('decision','予算内（実測要）')]:
        address=manifest['output_cells'][key].replace('$','')
        require(cached['21_秒数試算'][address].value==expected,f"Excel初期結果 {key}={expected}")
    require(len(wb['21_秒数試算'].data_validations.dataValidation)==35,"35入力に値域検証あり")
    require(len(wb['22_時間軸図']._charts)==2,"時間軸グラフ2個")
    with ZipFile(REPORT) as archive:
        names=archive.namelist()
        require(not any('vbaProject' in n or 'externalLinks' in n for n in names),"マクロ・外部数式リンクなし")
        ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        styles=ET.fromstring(archive.read('xl/styles.xml'))
        colors=[]
        for item in styles.iter():
            color=item.attrib.get('rgb')
            if color and len(color)==8:
                colors.append(color)
        require(all(c[2:4]==c[4:6]==c[6:8] for c in colors),"セルのRGB指定は全てモノトーン")
    summary={"status":"passed","check_count":len(checks),"sheets":len(wb.sheetnames),"formula_count":formulas,"font_names":sorted(fonts),"json_examples":len(parsed),"markdown_characters":len(md),"xlsx_bytes":REPORT.stat().st_size,"markdown_bytes":MARKDOWN.stat().st_size,"unverified":["AWS上の受理・配備・起動","実イメージと専用プローブ","EAPとreadiness","実測していない時間予算","Excel本体での表示・再計算（自動操作の起動が無応答）"],"checks":checks}
    (OUT/'validation.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({k:v for k,v in summary.items() if k!='checks'},ensure_ascii=False))


if __name__=='__main__':
    validate()
